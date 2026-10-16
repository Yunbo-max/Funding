#!/usr/bin/env python3
"""Conservative local helpers for Funding Autopilot (no network or LLM calls)."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import urlparse

VERSION = '0.1.0'
REQUIRED_EVIDENCE = {'scope', 'eligibility', 'deadline'}
ROUTES = ('independent', 'new_pi')


def read_json(path):
    with Path(path).open('r', encoding='utf-8') as fh:
        return json.load(fh)


def save_json(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def utc_now():
    return datetime.now(timezone.utc)


def parse_datetime(value):
    """Require offset-aware timestamps to avoid silently guessing call time zones."""
    if not isinstance(value, str) or not value.strip():
        raise ValueError('missing ISO datetime with UTC offset')
    dt = datetime.fromisoformat(value.replace('Z', '+00:00'))
    if dt.tzinfo is None or dt.utcoffset() is None:
        raise ValueError('datetime must include timezone offset')
    return dt.astimezone(timezone.utc)


def parse_date(value):
    if not isinstance(value, str) or not re.fullmatch(r'\d{4}-\d{2}-\d{2}', value):
        raise ValueError('date must be YYYY-MM-DD')
    return date.fromisoformat(value)


def extract_cv(path):
    """Extract readable text; no user data leaves the current filesystem."""
    path = Path(path)
    if not path.is_file():
        raise ValueError(f'CV not found: {path}')
    suffix = path.suffix.lower()
    if suffix in ('.txt', '.md'):
        text = path.read_text(encoding='utf-8-sig')
    elif suffix == '.pdf':
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ValueError("PDF support missing: pip install pypdf") from exc
        reader = PdfReader(str(path))
        text = '\n\n'.join((page.extract_text() or '') for page in reader.pages)
    elif suffix == '.docx':
        try:
            from docx import Document
        except ImportError as exc:
            raise ValueError("DOCX support missing: pip install python-docx") from exc
        doc = Document(str(path))
        text = '\n'.join([p.text for p in doc.paragraphs] +
                         [cell.text for table in doc.tables for row in table.rows for cell in row.cells])
    else:
        raise ValueError('Supported CV file types: PDF, DOCX, MD and TXT')
    if len(text.strip()) < 40:
        raise ValueError('Insufficient extractable text; check file or scanned/OCR-only PDF manually')
    return text


def intake(cv, output):
    cv = Path(cv)
    text = extract_cv(cv)
    output = Path(output)
    extracted = output.with_suffix('.cv.txt')
    extracted.parent.mkdir(parents=True, exist_ok=True)
    extracted.write_text(text, encoding='utf-8')
    profile = {
        'schema_version': VERSION,
        'cv_file': cv.name,
        'cv_sha256': hashlib.sha256(cv.read_bytes()).hexdigest(),
        'extracted_text_file': str(extracted),
        'text_characters': len(text),
        'fields_verified': False,
        'fact_evidence': [],
        'research_terms': [],
        'residence_country': None,
        'phd_submitted_on': None,
        'expected_phd_submission_on': None,
        'phd_awarded_on': None,
        'expected_phd_award_on': None,
        'host_confirmed': False,
        'new_pi_confirmed': False,
        'preferred_routes': list(ROUTES),
    }
    save_json(output, profile)
    return profile


def call_source_problems(call, as_of):
    problems = []
    url = call.get('official_url') or ''
    parsed = urlparse(url)
    if parsed.scheme != 'https' or not parsed.netloc:
        problems.append('missing HTTPS official call URL')
    try:
        checked = parse_datetime(call.get('verified_at'))
        if checked > as_of + timedelta(minutes=5):
            problems.append('checked-at timestamp is in the future')
        elif as_of - checked > timedelta(days=30):
            problems.append('official source check is older than 30 days; re-open current call')
    except (ValueError, TypeError):
        problems.append('official verified_at with UTC offset is missing/invalid')
    evidence = call.get('evidence') or []
    fields = {e.get('field') for e in evidence if isinstance(e, dict) and e.get('url') and e.get('locator')}
    for field in sorted(REQUIRED_EVIDENCE - fields):
        problems.append(f'missing primary source evidence: {field}')
    if not call.get('id') or not call.get('round'):
        problems.append('call ID and explicit round are required')
    if not isinstance(call.get('routes'), list) or not call.get('routes'):
        problems.append('applicant routes are not documented')
    status = call.get('call_status')
    if status not in ('open', 'rolling', 'closed'):
        problems.append('call open/rolling/closed status not verified')
    dls = call.get('deadlines') or {}
    if status != 'rolling':
        try:
            parse_datetime(dls.get('funder'))
        except (ValueError, TypeError):
            problems.append('funder deadline missing/invalid (must include offset)')
    req = call.get('requirements') or {}
    internal_status = req.get('internal_deadline_status')
    if internal_status not in ('verified', 'confirmed_none', 'unknown'):
        problems.append('internal deadline status is not recorded')
    if internal_status == 'verified':
        try:
            parse_datetime(dls.get('internal'))
        except (ValueError, TypeError):
            problems.append('verified internal deadline missing/invalid')
    return problems


def token_set(values):
    return {token for s in values if isinstance(s, str)
            for token in re.findall(r'[a-z0-9]+', s.lower()) if len(token) > 2}


def fit_score(profile, call, route, as_of):
    """Only a transparent heuristic of fit; never funding/acceptance probability."""
    applicant = token_set(profile.get('research_terms') or [])
    target = token_set(call.get('scope_terms') or [])
    overlap = len(applicant & target) / max(1, len(target))
    topic = round(45 * overlap)
    independence = 20 if route == 'independent' else 8
    salary = 20 if call.get('salary_support') == 'yes' else 0
    dls = call.get('deadlines') or {}
    try:
        delta = (parse_datetime(dls['funder']) - as_of).days
        readiness = 15 if delta >= 30 else (8 if delta >= 7 else 2)
    except (ValueError, TypeError, KeyError):
        readiness = 5 if call.get('call_status') == 'rolling' else 0
    return min(100, topic + independence + salary + readiness)


def assess(profile, call, route, as_of=None):
    if as_of is None:
        as_of = utc_now()
    result = {'id': call.get('id'), 'title': call.get('title'), 'funder': call.get('funder'),
              'round': call.get('round'), 'route': route, 'official_url': call.get('official_url'),
              'verified_at': call.get('verified_at'), 'status': None, 'reasons': [],
              'fit_score_not_probability': None}
    if not profile.get('cv_sha256'):
        result.update(status='unverified', reasons=['personalized checks require uploaded CV hash'])
        return result
    source_gaps = call_source_problems(call, as_of)
    if source_gaps:
        result.update(status='unverified', reasons=source_gaps)
        return result
    deadlines = call.get('deadlines') or {}
    if call['call_status'] == 'closed':
        result.update(status='expired', reasons=['official call marked closed'])
        return result
    for name in ('internal', 'funder'):
        when = deadlines.get(name)
        if when and parse_datetime(when) < as_of:
            result.update(status='expired', reasons=[f'{name} deadline has passed'])
            return result
    fail, uncertain = [], []
    if profile.get('fields_verified') is not True:
        uncertain.append('CV-derived facts have not been reviewed against the uploaded CV')
    if route not in (call.get('routes') or []):
        fail.append(f'{route} application route not permitted for this round')
    req = call.get('requirements') or {}
    for milestone, profile_key in (('requires_phd_submission_by', 'phd_submitted_on'),
                                   ('requires_phd_award_by', 'phd_awarded_on')):
        limit = req.get(milestone)
        if limit:
            try:
                cutoff = parse_date(limit)
                actual = profile.get(profile_key)
                if not actual:
                    uncertain.append(f'{profile_key} not yet evidenced (expected date is not proof)')
                elif parse_date(actual) > as_of.date():
                    uncertain.append(f'{profile_key} is in the future and cannot be treated as completed')
                elif parse_date(actual) > cutoff:
                    fail.append(f'{profile_key} is after permitted {cutoff.isoformat()}')
            except (ValueError, TypeError):
                uncertain.append(f'{milestone} or {profile_key} contains an invalid date')
    eligible_countries = req.get('eligible_residence_countries') or []
    if eligible_countries:
        country = profile.get('residence_country')
        if not country:
            uncertain.append('residence_country must be confirmed for this call')
        elif country not in eligible_countries:
            fail.append('residence_country not in verified eligible countries')
    if req.get('requires_host') is True and profile.get('host_confirmed') is not True:
        uncertain.append('host acceptance/eligibility is not confirmed')
    if (route == 'new_pi' or req.get('requires_new_pi') is True) and profile.get('new_pi_confirmed') is not True:
        uncertain.append('new PI acceptance and eligibility are not confirmed')
    if req.get('internal_deadline_status') == 'unknown':
        uncertain.append('receiving institution internal deadline remains unknown')
    manual = req.get('manual_gate_status')
    if manual == 'fail':
        fail.append('one or more additional manual eligibility conditions failed')
    elif manual != 'pass':
        uncertain.append('other call-specific eligibility conditions have not been signed off')
    result['fit_score_not_probability'] = fit_score(profile, call, route, as_of)
    if fail:
        result.update(status='ineligible', reasons=fail + uncertain)
    elif uncertain:
        result.update(status='conditional', reasons=uncertain)
    else:
        result.update(status='eligible', reasons=['encoded eligibility and deadlines pass; independent institutional review still required'])
    return result


def shortlist(profile, calls, as_of=None):
    if as_of is None:
        as_of = utc_now()
    if not profile.get('cv_sha256'):
        raise ValueError('CV upload/intake required: missing cv_sha256')
    if not isinstance(calls, list):
        raise ValueError('calls JSON must be an array')
    routes = profile.get('preferred_routes') or list(ROUTES)
    if not set(routes).issubset(ROUTES):
        raise ValueError('preferred_routes must use independent and/or new_pi')
    records = [assess(profile, call, route, as_of) for call in calls for route in routes]
    order = {'eligible': 0, 'conditional': 1, 'unverified': 2, 'ineligible': 3, 'expired': 4}
    records.sort(key=lambda r: (order.get(r['status'], 9),
                                -(r['fit_score_not_probability'] or 0),
                                str(r.get('id')), r['route']))
    return {'schema_version': VERSION, 'checked_at_utc': as_of.isoformat(),
            'warning': 'Scores are fit heuristics, NOT funding success probabilities. Human and institution verification remain necessary.',
            'results': records}


def prepare(call, output, as_of=None):
    """Prepare a funder-section-shaped blank draft; the AI host writes the prose."""
    if as_of is None:
        as_of = utc_now()
    problems = call_source_problems(call, as_of)
    if problems:
        raise ValueError('cannot prepare official-shaped draft before source verification: ' + '; '.join(problems))
    specs = call.get('sections') or []
    if not specs:
        raise ValueError("cannot prepare draft without this call's actual section specifications")
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    title = str(call.get('title') or call['id']).replace('\n', ' ').strip()
    lines = [f'# {title}', '', '**DRAFT — HUMAN/PI/HOST REVIEW REQUIRED; NOT SUBMITTED**', '',
             f"Funder: {call.get('funder', '[TO CONFIRM]')}",
             f"Round: {call.get('round')}",
             f"Official call: {call.get('official_url')}",
             f"Call checked: {call.get('verified_at')}",
             f"Funder deadline: {(call.get('deadlines') or {}).get('funder') or 'rolling (verify current status)'}", '',
             '## Eligibility and institutional checks', '',
             '- [ ] Personal eligibility verified against current call guidance',
             '- [ ] Host and/or new PI acceptance verified if required',
             '- [ ] Internal submission deadline and budget rules checked',
             '- [ ] Institution and applicant approve the final submission', '']
    for spec in specs:
        name = str(spec.get('id') or 'unnamed-section')
        word_cap = spec.get('max_words')
        lines.extend([f'## {name}', '',
                      f"Word cap: {word_cap if word_cap is not None else '[TO CONFIRM]'}; required: {bool(spec.get('required', True))}",
                      '', '[EVIDENCE NEEDED: draft this section using verified CV and official call criteria]', ''])
    lines += ['## Mandatory documents', '']
    for document in call.get('required_documents') or []:
        lines.append(f'- [ ] {document}')
    lines += ['', '## Research-claim and source ledger', '',
              '| Claim | CV/publication/call evidence | Status |', '|---|---|---|',
              '| [TO CONFIRM] | [EVIDENCE NEEDED] | unresolved |', '']
    output.write_text('\n'.join(lines), encoding='utf-8')
    return {'file': str(output), 'call_id': call['id'], 'section_count': len(specs),
            'status': 'blank_draft_human_review_required'}


def count_words(text):
    return len(re.findall(r"\b[\w]+(?:['-][\w]+)*\b", text, flags=re.UNICODE))


def audit(proposal, call, as_of=None):
    if as_of is None:
        as_of = utc_now()
    issues = []
    if proposal.get('call_id') != call.get('id'):
        issues.append('proposal call_id does not match selected official call')
    issues.extend(call_source_problems(call, as_of))
    sections = proposal.get('sections') or {}
    if not call.get('sections'):
        issues.append('exact funder section specifications not provided')
    for spec in call.get('sections') or []:
        section_id = spec.get('id')
        content = sections.get(section_id, '')
        if spec.get('required', True) and (not isinstance(content, str) or not content.strip()):
            issues.append(f'required section missing: {section_id}')
        if isinstance(content, str) and isinstance(spec.get('max_words'), int):
            n = count_words(content)
            if n > spec['max_words']:
                issues.append(f'section {section_id} over word limit: {n} > {spec["max_words"]}')
    for key, text in sections.items():
        if not isinstance(text, str):
            issues.append(f'section {key} is not text')
        elif re.search(r'\[(?:TO CONFIRM|EVIDENCE NEEDED|HOST APPROVAL REQUIRED)[^\]]*\]', text, flags=re.I):
            issues.append(f'placeholder remains in section: {key}')
    missing_docs = set(call.get('required_documents') or []) - set(proposal.get('documents') or [])
    if missing_docs:
        issues.append('missing required documents: ' + ', '.join(sorted(missing_docs)))
    sources = {r.get('id'): r for r in proposal.get('evidence_index') or [] if isinstance(r, dict) and r.get('id')}
    claims = proposal.get('claims') or []
    if not claims:
        issues.append('claim-to-evidence ledger is empty')
    for claim in claims:
        claim_id = claim.get('id', 'unknown')
        if not claim.get('text'):
            issues.append(f'claim {claim_id} has no text')
        refs = claim.get('evidence_ids') or []
        if not refs:
            issues.append(f'claim {claim_id} has no evidence IDs')
        for ref in refs:
            evidence = sources.get(ref)
            if not evidence or evidence.get('verified') is not True or not evidence.get('locator'):
                issues.append(f'claim {claim_id} has missing/unverified evidence: {ref}')
    return {'schema_version': VERSION, 'audited_at_utc': as_of.isoformat(),
            'status': 'needs_repair' if issues else 'structured_draft_passed_human_review_required',
            'issues': issues,
            'human_review_required': True,
            'warning': 'This is a partial structure audit, NOT official policy verification or permission to submit.'}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='command', required=True)
    p = sub.add_parser('intake', help='extract user-uploaded CV into a private conservative profile')
    p.add_argument('--cv', required=True)
    p.add_argument('--out', required=True)
    p = sub.add_parser('shortlist', help='classify agent-verified structured calls by route')
    p.add_argument('--profile', required=True)
    p.add_argument('--calls', required=True)
    p.add_argument('--out')
    p.add_argument('--as-of', help='aware ISO UTC/offset timestamp for repeatable checks')
    p = sub.add_parser('prepare', help='create blank draft with official-call section headings')
    p.add_argument('--call', required=True)
    p.add_argument('--out', required=True)
    p.add_argument('--as-of', help='aware ISO UTC/offset timestamp')
    p = sub.add_parser('audit', help='audit proposal structure against structured call spec')
    p.add_argument('--proposal', required=True)
    p.add_argument('--call', required=True)
    p.add_argument('--out')
    p.add_argument('--as-of', help='aware ISO UTC/offset timestamp')
    args = parser.parse_args(argv)
    try:
        if args.command == 'intake':
            data = intake(args.cv, args.out)
            summary = {'cv_sha256': data['cv_sha256'], 'characters': data['text_characters'],
                       'profile_file': args.out, 'fields_verified': False}
        elif args.command == 'shortlist':
            now = parse_datetime(args.as_of) if args.as_of else utc_now()
            data = shortlist(read_json(args.profile), read_json(args.calls), now)
            if args.out:
                save_json(args.out, data)
            summary = data
        elif args.command == 'prepare':
            now = parse_datetime(args.as_of) if args.as_of else utc_now()
            summary = prepare(read_json(args.call), args.out, now)
        else:
            now = parse_datetime(args.as_of) if args.as_of else utc_now()
            data = audit(read_json(args.proposal), read_json(args.call), now)
            if args.out:
                save_json(args.out, data)
            summary = data
        print(json.dumps(summary, ensure_ascii=False, indent=2))
        return 0
    except (OSError, ValueError, TypeError, json.JSONDecodeError) as exc:
        print(f'funding-autopilot: {exc}', file=sys.stderr)
        return 2


if __name__ == '__main__':
    sys.exit(main())
